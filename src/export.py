"""Excel export for research results."""

import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference, PieChart
from openpyxl.utils import get_column_letter

from .engine import SimulationResult
from .analysis import OverallStats


# Style constants
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
ACCENT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
GOOD_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
BAD_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _style_header_row(ws, row: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER


def _auto_width(ws):
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)


def export_to_excel(results: list[SimulationResult],
                    stats: OverallStats,
                    output_dir: str = "output") -> str:
    """Export research results to a formatted Excel workbook. Returns file path."""
    os.makedirs(output_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{stats.company.replace(' ', '_')}_research_{ts}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()

    # ── Sheet 1: Executive Summary ──
    ws_summary = wb.active
    ws_summary.title = "Executive Summary"
    ws_summary.sheet_properties.tabColor = "2F5496"

    ws_summary.merge_cells("A1:D1")
    title_cell = ws_summary["A1"]
    title_cell.value = f"LLM Visibility Report — {stats.company}"
    title_cell.font = Font(bold=True, size=16, color="2F5496")

    ws_summary["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws_summary["A2"].font = Font(italic=True, color="666666")

    row = 4
    metrics = [
        ("Total Simulations", stats.total_simulations),
        ("Overall Mention Rate", f"{stats.overall_mention_rate:.1%}"),
        ("Average Sentiment", f"{stats.overall_avg_sentiment:+.2f}"),
        ("Average Ranking Position", stats.overall_avg_position or "N/A"),
        ("LLM Visibility Score", f"{stats.visibility_score:.1f} / 100"),
    ]
    for label, value in metrics:
        ws_summary.cell(row=row, column=1, value=label).font = Font(bold=True)
        val_cell = ws_summary.cell(row=row, column=2, value=str(value))
        if label == "LLM Visibility Score":
            val_cell.font = Font(bold=True, size=14, color="2F5496")
        row += 1

    # Competitor summary
    if stats.competitor_summary:
        row += 1
        ws_summary.cell(row=row, column=1, value="Competitor Mention Rates").font = Font(bold=True, size=12)
        row += 1
        headers = ["Company", "Mention Rate"]
        for c, h in enumerate(headers, 1):
            ws_summary.cell(row=row, column=c, value=h)
        _style_header_row(ws_summary, row, len(headers))
        row += 1
        # Target company first
        ws_summary.cell(row=row, column=1, value=f"{stats.company} (target)")
        ws_summary.cell(row=row, column=2, value=f"{stats.overall_mention_rate:.1%}")
        ws_summary.cell(row=row, column=1).fill = ACCENT_FILL
        ws_summary.cell(row=row, column=2).fill = ACCENT_FILL
        row += 1
        for comp, rate in stats.competitor_summary.items():
            ws_summary.cell(row=row, column=1, value=comp)
            ws_summary.cell(row=row, column=2, value=f"{rate:.1%}")
            row += 1

    _auto_width(ws_summary)

    # ── Sheet 2: Category Breakdown ──
    ws_cat = wb.create_sheet("Category Breakdown")
    ws_cat.sheet_properties.tabColor = "548235"

    headers = ["Category", "Runs", "Mentions", "Mention Rate",
               "Avg Position", "Avg Sentiment", "Sentiment StdDev"]
    for c, h in enumerate(headers, 1):
        ws_cat.cell(row=1, column=c, value=h)
    _style_header_row(ws_cat, 1, len(headers))

    for i, cs in enumerate(stats.category_stats, 2):
        ws_cat.cell(row=i, column=1, value=cs.category.replace("_", " ").title())
        ws_cat.cell(row=i, column=2, value=cs.total_runs)
        ws_cat.cell(row=i, column=3, value=cs.mention_count)
        ws_cat.cell(row=i, column=4, value=f"{cs.mention_rate:.1%}")
        ws_cat.cell(row=i, column=5, value=cs.avg_position or "N/A")
        ws_cat.cell(row=i, column=6, value=f"{cs.avg_sentiment:+.3f}")
        ws_cat.cell(row=i, column=7, value=f"{cs.sentiment_std:.3f}")
        # Colour code mention rate
        rate_cell = ws_cat.cell(row=i, column=4)
        rate_cell.fill = GOOD_FILL if cs.mention_rate >= 0.6 else BAD_FILL if cs.mention_rate < 0.3 else PatternFill()

    _auto_width(ws_cat)

    # Bar chart: mention rate by category
    if stats.category_stats:
        chart = BarChart()
        chart.type = "col"
        chart.title = "Mention Rate by Category"
        chart.y_axis.title = "Mention Rate"
        chart.x_axis.title = "Category"
        # Add data for chart (numeric mention rates)
        chart_row = len(stats.category_stats) + 3
        for j, cs in enumerate(stats.category_stats):
            ws_cat.cell(row=chart_row + j, column=1, value=cs.category.replace("_", " ").title())
            ws_cat.cell(row=chart_row + j, column=2, value=cs.mention_rate)
        data = Reference(ws_cat, min_col=2, min_row=chart_row, max_row=chart_row + len(stats.category_stats) - 1)
        cats = Reference(ws_cat, min_col=1, min_row=chart_row, max_row=chart_row + len(stats.category_stats) - 1)
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        chart.shape = 4
        chart.width = 18
        chart.height = 10
        ws_cat.add_chart(chart, f"A{chart_row + len(stats.category_stats) + 1}")

    # ── Sheet 3: Raw Results ──
    ws_raw = wb.create_sheet("Raw Results")
    ws_raw.sheet_properties.tabColor = "BF8F00"

    raw_headers = ["#", "Category", "Prompt", "Mentioned", "Position",
                   "Sentiment", "Response (truncated)", "Latency (ms)"]
    for c, h in enumerate(raw_headers, 1):
        ws_raw.cell(row=1, column=c, value=h)
    _style_header_row(ws_raw, 1, len(raw_headers))

    for i, r in enumerate(results, 2):
        ws_raw.cell(row=i, column=1, value=i - 1)
        ws_raw.cell(row=i, column=2, value=r.category.replace("_", " ").title())
        ws_raw.cell(row=i, column=3, value=r.prompt[:120])
        ws_raw.cell(row=i, column=4, value="Yes" if r.company_mentioned else "No")
        ws_raw.cell(row=i, column=5, value=r.mention_position or "—")
        ws_raw.cell(row=i, column=6, value=f"{r.sentiment_score:+.3f}")
        ws_raw.cell(row=i, column=7, value=r.response[:300])
        ws_raw.cell(row=i, column=8, value=r.latency_ms)
        # Colour code mention
        mention_cell = ws_raw.cell(row=i, column=4)
        mention_cell.fill = GOOD_FILL if r.company_mentioned else BAD_FILL

    _auto_width(ws_raw)

    wb.save(filepath)
    return filepath
